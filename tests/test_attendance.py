import unittest
from datetime import time
from io import BytesIO

from openpyxl import Workbook, load_workbook

from src.attendance_logic import (
    buat_rekap_bulanan,
    buat_rekap_semua_karyawan,
    hitung_pulang_awal,
    hitung_terlambat,
    ringkas_rekap,
)
from src.attendance_parser import baca_file_absensi, normalisasi_jam
from src.audit_export import buat_excel_audit_individu, buat_excel_audit_semua
from src.audit_logic import buat_ringkasan_audit, rincian_karyawan_audit, ringkas_audit
from src.audit_parser import baca_file_audit
from src.excel_export import buat_excel_rekap, buat_excel_rekap_semua
from src.file_parser import deteksi_jenis_file, proses_file


def data_uji():
    return {
        "karyawan": {
            "K001": {"id": "K001", "nama": "Budi", "departemen": "Produksi"},
            "K002": {"id": "K002", "nama": "Siti", "departemen": "Packing"},
        },
        "data_harian": {
            1: {
                "K001": {
                    "pagi_masuk": "07:45",
                    "pagi_pulang": "12:00",
                    "siang_masuk": "13:00",
                    "siang_pulang": "16:00",
                    "lembur_masuk": "17:00",
                    "lembur_pulang": "19:00",
                },
                "K002": {
                    "pagi_masuk": "07:25",
                    "pagi_pulang": "12:00",
                    "siang_masuk": "13:00",
                    "siang_pulang": "16:30",
                    "lembur_masuk": "",
                    "lembur_pulang": "",
                },
            },
            2: {
                "K001": {
                    "pagi_masuk": "07:30",
                    "pagi_pulang": "12:00",
                    "siang_masuk": "",
                    "siang_pulang": "",
                    "lembur_masuk": "",
                    "lembur_pulang": "",
                }
            },
        },
    }


class TestLogikaAbsensi(unittest.TestCase):
    def test_toleransi_adalah_ambang(self):
        self.assertEqual(hitung_terlambat("07:35", "07:30", 5), 0)
        self.assertEqual(hitung_terlambat("07:36", "07:30", 5), 6)
        self.assertEqual(hitung_pulang_awal("16:25", "16:30", 5), 0)
        self.assertEqual(hitung_pulang_awal("16:24", "16:30", 5), 6)

    def test_status_durasi_dan_hari_kerja(self):
        hasil = buat_rekap_bulanan(
            data_uji(), "K001", 2026, 9, hari_kerja=[0, 1, 2, 3, 4]
        )
        hari_pertama = hasil[0]
        self.assertEqual(hari_pertama["Tanggal"].day, 1)
        self.assertEqual(hari_pertama["Status"], "TERLAMBAT & PULANG AWAL")
        self.assertEqual(hari_pertama["Jam Kerja"], "7:15")
        self.assertEqual(hari_pertama["Jam Lembur"], "2:00")
        self.assertEqual(hasil[1]["Status"], "SCAN TIDAK LENGKAP")
        self.assertTrue(all(item["Tanggal"].weekday() < 5 for item in hasil))

    def test_rekap_semua_dan_ringkasan(self):
        hasil = buat_rekap_semua_karyawan(
            data_uji(), ["K001", "K002"], 2026, 9, hari_kerja=[1]
        )
        self.assertEqual(len(hasil), 10)
        self.assertTrue(all(item["ID"] in {"K001", "K002"} for item in hasil))
        total = ringkas_rekap(hasil)
        self.assertEqual(total["hadir"], 2)
        self.assertEqual(total["terlambat"], 1)
        self.assertEqual(total["pulang_awal"], 1)
        self.assertEqual(total["total_menit_lembur"], 120)


class TestParser(unittest.TestCase):
    def test_normalisasi_variasi_jam(self):
        self.assertEqual(normalisasi_jam(time(7, 30)), "07:30")
        self.assertEqual(normalisasi_jam(0.5), "12:00")
        self.assertEqual(normalisasi_jam("07:30:00"), "07:30")
        self.assertEqual(normalisasi_jam(305), "")

    def test_membaca_format_mesin(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "1"
        ws.cell(9, 3, "K001")
        ws.cell(9, 4, "Produksi")
        ws.cell(9, 5, "Budi")
        ws.cell(9, 6, time(7, 30))
        ws.cell(9, 7, time(12, 0))
        ws.cell(9, 8, time(13, 0))
        ws.cell(9, 9, time(16, 30))
        ws.cell(9, 10, 305)
        ws.cell(9, 11, 305)
        wb.create_sheet("Ringkasan")
        output = BytesIO()
        wb.save(output)

        parsed = baca_file_absensi(output.getvalue())
        self.assertEqual(parsed["karyawan"]["K001"]["nama"], "Budi")
        self.assertEqual(parsed["data_harian"][1]["K001"]["pagi_masuk"], "07:30")
        self.assertEqual(parsed["data_harian"][1]["K001"]["lembur_masuk"], "")


class TestExcelExport(unittest.TestCase):
    def test_export_individu_dan_semua(self):
        parsed = data_uji()
        rincian = buat_rekap_bulanan(parsed, "K001", 2026, 9, hari_kerja=[1])
        file_individu = buat_excel_rekap(rincian, "Budi", "K001", "Produksi")
        wb_individu = load_workbook(file_individu, data_only=True)
        self.assertEqual(wb_individu.sheetnames, ["Absensi"])
        self.assertEqual(wb_individu["Absensi"].max_column, 17)
        self.assertEqual(wb_individu["Absensi"]["C2"].value, "K001")

        ringkasan = [
            {
                "No": 1,
                "ID": "K001",
                "Nama": "Budi",
                "Departemen": "Produksi",
                "Hadir": 1,
                "Terlambat": 1,
                "Pulang Awal": 1,
                "Scan Tidak Lengkap": 0,
                "Tidak Hadir": 0,
                "Tidak Ada Data": 4,
                "Total Jam Kerja": "7:15",
                "Total Jam Lembur": "2:00",
            }
        ]
        file_semua = buat_excel_rekap_semua(ringkasan, rincian)
        wb_semua = load_workbook(file_semua, data_only=True)
        self.assertEqual(wb_semua.sheetnames, ["Ringkasan", "Rincian"])
        self.assertEqual(wb_semua["Ringkasan"]["B2"].value, "K001")


def file_audit_uji():
    wb = Workbook()
    ws = wb.active
    ws.title = "PRODUKSI"
    ws["B1"] = "MONTH : 2026/06"
    ws["N3"] = 1
    ws["N4"] = "jam normal"
    ws["O4"] = "um"
    ws["P4"] = "shift malam"
    ws["Q4"] = "jam sakit"
    ws["R4"] = "cuti tahunan"
    ws["S4"] = "cuti khusus"
    ws["T4"] = "jam lembur"
    ws["U4"] = "um lembur"
    ws["V3"] = 2
    ws["V4"] = "jam normal"
    ws["A5"] = 1
    ws["B5"] = 761001
    ws["C5"] = "Produksi"
    ws["D5"] = "Operator"
    ws["E5"] = "Budi"
    ws["N5"] = 7
    ws["O5"] = 1
    ws["T5"] = 2
    ws["V5"] = 7
    ws["JE4"] = "LIBUR YANG DIBAYAR"
    ws["JE5"] = "=N5"
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


class TestAudit(unittest.TestCase):
    def test_deteksi_dan_parser_audit(self):
        file_bytes = file_audit_uji()
        self.assertEqual(deteksi_jenis_file(file_bytes), "audit")

        parsed = proses_file(file_bytes)
        self.assertEqual(parsed["periode"], {"tahun": 2026, "bulan": 6})
        self.assertEqual(parsed["karyawan"]["761001"]["nama"], "Budi")
        self.assertEqual(len(parsed["rincian"]), 2)
        self.assertEqual(parsed["rincian"][0]["status"], "LIBUR DIBAYAR")
        self.assertEqual(parsed["rincian"][1]["status"], "HADIR")
        self.assertEqual(parsed["rincian"][0]["jam_lembur"], 2)

    def test_ringkasan_dan_export_audit(self):
        parsed = baca_file_audit(file_audit_uji())
        rincian = rincian_karyawan_audit(parsed, "761001")
        total = ringkas_audit(rincian)
        self.assertEqual(total["hari_hadir"], 1)
        self.assertEqual(total["jam_normal"], 14)
        self.assertEqual(total["jam_hadir"], 7)
        self.assertEqual(total["libur_dibayar"], 7)
        self.assertEqual(total["jam_lembur"], 2)

        ringkasan = buat_ringkasan_audit(parsed, ["761001"])
        self.assertEqual(ringkasan[0]["Jam Hadir"], 7)

        individu = buat_excel_audit_individu(rincian)
        wb_individu = load_workbook(individu, data_only=True)
        self.assertEqual(wb_individu.sheetnames, ["Rincian Audit"])
        self.assertEqual(wb_individu["Rincian Audit"]["B2"].value, "761001")

        semua = buat_excel_audit_semua(ringkasan, rincian)
        wb_semua = load_workbook(semua, data_only=True)
        self.assertEqual(wb_semua.sheetnames, ["Ringkasan", "Rincian"])
        self.assertEqual(wb_semua["Ringkasan"]["B2"].value, "761001")


if __name__ == "__main__":
    unittest.main()

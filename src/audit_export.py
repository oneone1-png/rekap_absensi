"""Ekspor rekap audit ke Excel."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from src.audit_logic import rincian_untuk_tampilan


WARNA_HEADER = "1F4E78"
GARIS = Side(style="thin", color="B7C3D0")
BORDER = Border(left=GARIS, right=GARIS, top=GARIS, bottom=GARIS)


def _isi_sheet(sheet, rows):
    if not rows:
        sheet.append(["Tidak ada data"])
        return

    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])

    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=WARNA_HEADER)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(vertical="center")
            cell.border = BORDER

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    for column in sheet.columns:
        letter = column[0].column_letter
        panjang = max(len(str(cell.value or "")) for cell in column)
        sheet.column_dimensions[letter].width = min(max(panjang + 2, 11), 30)

    for cell in sheet["A"][1:]:
        if hasattr(cell.value, "year"):
            cell.number_format = "dd/mm/yyyy"


def _simpan(workbook):
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def buat_excel_audit_individu(rincian):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rincian Audit"
    _isi_sheet(sheet, rincian_untuk_tampilan(rincian))
    return _simpan(workbook)


def buat_excel_audit_semua(ringkasan, rincian):
    workbook = Workbook()
    sheet_ringkasan = workbook.active
    sheet_ringkasan.title = "Ringkasan"
    _isi_sheet(sheet_ringkasan, ringkasan)

    sheet_rincian = workbook.create_sheet("Rincian")
    _isi_sheet(sheet_rincian, rincian_untuk_tampilan(rincian))
    return _simpan(workbook)

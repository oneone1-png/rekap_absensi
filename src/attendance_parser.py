from io import BytesIO
from datetime import datetime, time, timedelta
from openpyxl import load_workbook


def normalisasi_id(value):
    if value is None:
        return ""

    if isinstance(value, (int, float)):
        return str(int(value))

    return str(value).strip()


def normalisasi_jam(value):
    if value is None:
        return ""

    # Pada file attendance Anda, nilai 305 dianggap kosong
    if value == 305:
        return ""

    if isinstance(value, datetime):
        return value.strftime("%H:%M")

    if isinstance(value, time):
        return value.strftime("%H:%M")

    if isinstance(value, timedelta):
        total_menit = int(round(value.total_seconds() / 60)) % (24 * 60)
        return f"{total_menit // 60:02d}:{total_menit % 60:02d}"

    # Excel kadang menyimpan jam sebagai angka pecahan
    if isinstance(value, (int, float)) and 0 <= value < 1:
        total_menit = int(round(value * 24 * 60))

        jam = (total_menit // 60) % 24
        menit = total_menit % 60

        return f"{jam:02d}:{menit:02d}"

    text = str(value).strip()

    if text == "305":
        return ""

    # Terima nilai teks 07:30:00 atau tanggal+jam dari variasi ekspor mesin.
    for format_jam in ("%H:%M", "%H:%M:%S", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, format_jam).strftime("%H:%M")
        except ValueError:
            continue

    return text


def baca_file_absensi(file_bytes):

    if not file_bytes:
        raise ValueError("File attendance kosong.")

    workbook = load_workbook(
        BytesIO(file_bytes),
        read_only=True,
        data_only=True
    )

    data_harian = {}
    karyawan = {}

    # Baca seluruh worksheet
    for sheet in workbook.worksheets:

        # Nama sheet, misalnya 1, 2, 3 ... 31
        nama_sheet = sheet.title.strip()

        # Abaikan sheet selain angka
        if not nama_sheet.isdigit():
            continue

        # Nama sheet digunakan sebagai tanggal
        tanggal = int(nama_sheet)

        # Hanya menerima tanggal 1 sampai 31
        if tanggal < 1 or tanggal > 31:
            continue

        data_harian.setdefault(tanggal, {})

        # Data karyawan pada file Anda dimulai sekitar baris 9
        for row in sheet.iter_rows(
            min_row=9,
            max_col=11,
            values_only=True
        ):

            # C = ID
            id_karyawan = normalisasi_id(row[2])

            # D = Departemen
            departemen = row[3]

            # E = Nama
            nama = row[4]

            # Lewati baris yang bukan data karyawan
            if not id_karyawan or not nama:
                continue

            record = {
                "id": id_karyawan,
                "departemen": (
                    str(departemen).strip()
                    if departemen is not None
                    else ""
                ),
                "nama": str(nama).strip(),

                # F
                "pagi_masuk": normalisasi_jam(row[5]),

                # G
                "pagi_pulang": normalisasi_jam(row[6]),

                # H
                "siang_masuk": normalisasi_jam(row[7]),

                # I
                "siang_pulang": normalisasi_jam(row[8]),

                # J
                "lembur_masuk": normalisasi_jam(row[9]),

                # K
                "lembur_pulang": normalisasi_jam(row[10]),
            }

            # Simpan data karyawan berdasarkan tanggal
            data_harian[tanggal][id_karyawan] = record

            # Simpan master karyawan
            karyawan[id_karyawan] = {
                "id": id_karyawan,
                "nama": str(nama).strip(),
                "departemen": (
                    str(departemen).strip()
                    if departemen is not None
                    else ""
                )
            }

    workbook.close()

    return {
        "karyawan": karyawan,
        "data_harian": data_harian
    }

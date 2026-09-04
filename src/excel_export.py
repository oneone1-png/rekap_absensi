from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


DETAIL_HEADERS = [
    "NO\n序",
    "Tanggal\n日期",
    "ID\n員工號",
    "Departemen\n部門",
    "Nama\n姓名",
    "Pagi Masuk\n上午上班",
    "Pagi Pulang\n上午下班",
    "Siang Masuk\n下午上班",
    "Siang Pulang\n下午下班",
    "Lembur Masuk\n加班上班",
    "Lembur Pulang\n加班下班",
    "Jam Lembur\n加班時",
    "Jam Kerja\n工作時數",
    "Terlambat (menit)\n遲到",
    "Pulang Awal (menit)\n早退",
    "Status\n狀態",
    "Keterangan\n備註",
]

DETAIL_KEYS = [
    "No",
    "Tanggal",
    "ID",
    "Departemen",
    "Nama",
    "Pagi Masuk",
    "Pagi Pulang",
    "Siang Masuk",
    "Siang Pulang",
    "Lembur Masuk",
    "Lembur Pulang",
    "Jam Lembur",
    "Jam Kerja",
    "Terlambat",
    "Pulang Awal",
    "Status",
    "Keterangan",
]

RINGKASAN_HEADERS = [
    "No",
    "ID",
    "Nama",
    "Departemen",
    "Hadir",
    "Terlambat",
    "Pulang Awal",
    "Izin",
    "Sakit",
    "Cuti",
    "Libur Nasional",
    "Cuti Hamil",
    "Scan Tidak Lengkap",
    "Tidak Hadir",
    "Tidak Ada Data",
    "Total Jam Kerja",
    "Total Jam Lembur",
    

]


THIN = Side(border_style="thin", color="B7C3D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=9, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=9)
STATUS_FILLS = {
    "HADIR": "E2F0D9",
    "TERLAMBAT": "FFF2CC",
    "PULANG AWAL": "FCE4D6",
    "TERLAMBAT & PULANG AWAL": "F8CBAD",
    "SCAN TIDAK LENGKAP": "FFE699",
    "TIDAK HADIR": "F4B084",
    "TIDAK ADA DATA": "E7E6E6",
}


def _format_header(ws, jumlah_kolom):
    for cell in ws[1][:jumlah_kolom]:
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.row_dimensions[1].height = 48
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _format_body(ws, jumlah_kolom, kolom_nama=5):
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=jumlah_kolom):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = BODY_FONT
        row[kolom_nama - 1].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row[0].row].height = 21


def _atur_print(ws, landscape=True):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"
    ws.sheet_view.showGridLines = False


def _buat_sheet_detail(ws, hasil, identitas_tetap=None):
    ws.append(DETAIL_HEADERS)
    identitas_tetap = identitas_tetap or {}

    for index, data in enumerate(hasil, start=1):
        row = []
        for key in DETAIL_KEYS:
            value = data.get(key, "")
            if key in {"ID", "Departemen", "Nama"} and index == 1:
                value = identitas_tetap.get(key, value)
            row.append(value)
        ws.append(row)

    _format_header(ws, len(DETAIL_HEADERS))
    _format_body(ws, len(DETAIL_HEADERS))

    for row_number in range(2, ws.max_row + 1):
        ws.cell(row_number, 2).number_format = "yyyy-mm-dd"
        status_cell = ws.cell(row_number, 16)
        warna = STATUS_FILLS.get(str(status_cell.value), "")
        if warna:
            status_cell.fill = PatternFill("solid", fgColor=warna)

    widths = [5, 13, 11, 17, 26, 12, 12, 12, 12, 13, 13, 12, 12, 15, 17, 27, 30,]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width

    _atur_print(ws)


def _buat_sheet_ringkasan(ws, ringkasan):
    ws.append(RINGKASAN_HEADERS)
    for data in ringkasan:
        ws.append([data.get(header, "") for header in RINGKASAN_HEADERS])

    _format_header(ws, len(RINGKASAN_HEADERS))
    _format_body(ws, len(RINGKASAN_HEADERS), kolom_nama=3)
    widths = [5, 12, 27, 18, 10, 12, 13, 21, 14, 18, 18, 19]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    _atur_print(ws)


def _simpan(wb):
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def buat_excel_rekap(hasil, nama_karyawan, id_karyawan, departemen):
    wb = Workbook()
    ws = wb.active
    ws.title = "Absensi"
    _buat_sheet_detail(
        ws,
        hasil,
        {
            "ID": id_karyawan,
            "Departemen": departemen,
            "Nama": nama_karyawan,
        },
    )
    return _simpan(wb)


def buat_excel_rekap_semua(ringkasan, rincian):
    wb = Workbook()
    ws_ringkasan = wb.active
    ws_ringkasan.title = "Ringkasan"
    _buat_sheet_ringkasan(ws_ringkasan, ringkasan)

    ws_rincian = wb.create_sheet("Rincian")
    _buat_sheet_detail(ws_rincian, rincian)
    return _simpan(wb)

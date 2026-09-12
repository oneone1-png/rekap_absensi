"""Parser untuk workbook audit bulanan PRODUKSI dan OFFICE."""

import calendar
import re
from datetime import date
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from src.attendance_parser import normalisasi_id


SHEET_AUDIT = ("PRODUKSI", "OFFICE")
NAMA_FIELD_HARIAN = (
    "jam_normal",
    "uang_makan",
    "shift_malam",
    "jam_sakit",
    "cuti_tahunan",
    "cuti_khusus",
    "jam_lembur",
    "uang_makan_lembur",
)


def _angka(value):
    """Mengubah isi sel audit menjadi angka tanpa menghilangkan nilai desimal."""
    if value in (None, ""):
        return 0
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return value

    text = str(value).strip().replace(",", ".")
    try:
        number = float(text)
    except ValueError:
        return 0
    return int(number) if number.is_integer() else number


def _baca_periode(workbook):
    pola = re.compile(r"(20\d{2})\s*[/.-]\s*(\d{1,2})")
    for nama_sheet in SHEET_AUDIT:
        if nama_sheet not in workbook.sheetnames:
            continue
        sheet = workbook[nama_sheet]
        for row in sheet.iter_rows(min_row=1, max_row=4, max_col=8, values_only=True):
            for value in row:
                cocok = pola.search(str(value or ""))
                if cocok:
                    tahun, bulan = map(int, cocok.groups())
                    if 1 <= bulan <= 12:
                        return tahun, bulan
    raise ValueError(
        "Periode audit tidak ditemukan. Pastikan bagian atas sheet berisi "
        "teks seperti 'MONTH : 2026/06'."
    )


def _kolom_tanggal(sheet, jumlah_hari):
    """Mencari kolom awal tiap blok tanggal dari nomor pada baris 3."""
    hasil = []
    for column in range(1, sheet.max_column + 1):
        value = sheet.cell(3, column).value
        if isinstance(value, (int, float)) and int(value) == value:
            hari = int(value)
            if 1 <= hari <= jumlah_hari:
                hasil.append((hari, column))
    return hasil


def _status_harian(data):
    status = []
    if data["jam_hadir"] > 0:
        status.append("HADIR")
    if data["libur_dibayar"] > 0:
        status.append("LIBUR DIBAYAR")
    if data["jam_sakit"] > 0:
        status.append("SAKIT")
    if data["cuti_tahunan"] > 0:
        status.append("CUTI TAHUNAN")
    if data["cuti_khusus"] > 0:
        status.append("CUTI KHUSUS")
    if not status and any(data[field] > 0 for field in NAMA_FIELD_HARIAN):
        status.append("DATA LAIN")
    return " + ".join(status) if status else "TANPA DATA"


def _tanggal_libur_dibayar(sheet_nilai, sheet_formula, kolom_harian):
    """Membaca tanggal libur dari rumus total 'LIBUR YANG DIBAYAR'."""
    kolom_ke_hari = {column: hari for hari, column in kolom_harian}
    kolom_total = None
    for column in range(1, sheet_nilai.max_column + 1):
        header = str(sheet_nilai.cell(4, column).value or "").upper()
        if "LIBUR YANG DIBAYAR" in header:
            kolom_total = column
            break
    if kolom_total is None:
        return set()

    for row_number in range(5, sheet_nilai.max_row + 1):
        if not sheet_nilai.cell(row_number, 2).value:
            continue
        formula = sheet_formula.cell(row_number, kolom_total).value
        if not isinstance(formula, str) or not formula.startswith("="):
            continue
        hasil = set()
        for column_letter in re.findall(r"\$?([A-Z]{1,3})\$?\d+", formula.upper()):
            column = column_index_from_string(column_letter)
            if column in kolom_ke_hari:
                hasil.add(kolom_ke_hari[column])
        if hasil:
            return hasil
    return set()


def baca_file_audit(file_bytes):
    """Membaca file audit menjadi master karyawan dan rincian per tanggal."""
    if not file_bytes:
        raise ValueError("File audit kosong.")

    workbook = load_workbook(BytesIO(file_bytes), read_only=False, data_only=True)
    try:
        workbook_formula = load_workbook(
            BytesIO(file_bytes), read_only=False, data_only=False
        )
    except Exception:
        workbook.close()
        raise
    try:
        sheet_ditemukan = [name for name in SHEET_AUDIT if name in workbook.sheetnames]
        if not sheet_ditemukan:
            raise ValueError("Sheet PRODUKSI atau OFFICE tidak ditemukan.")

        tahun, bulan = _baca_periode(workbook)
        jumlah_hari = calendar.monthrange(tahun, bulan)[1]
        karyawan = {}
        rincian = []

        for nama_sheet in sheet_ditemukan:
            sheet = workbook[nama_sheet]
            kolom_harian = _kolom_tanggal(sheet, jumlah_hari)
            if not kolom_harian:
                continue
            tanggal_libur = _tanggal_libur_dibayar(
                sheet, workbook_formula[nama_sheet], kolom_harian
            )

            for row_number in range(5, sheet.max_row + 1):
                id_karyawan = normalisasi_id(sheet.cell(row_number, 2).value)
                nama = str(sheet.cell(row_number, 5).value or "").strip()
                if not id_karyawan or not nama:
                    continue

                departemen = str(sheet.cell(row_number, 3).value or "").strip()
                jabatan = str(sheet.cell(row_number, 4).value or "").strip()
                karyawan[id_karyawan] = {
                    "id": id_karyawan,
                    "nama": nama,
                    "departemen": departemen,
                    "jabatan": jabatan,
                    "unit": nama_sheet,
                }

                for hari, column_start in kolom_harian:
                    nilai = {
                        field: _angka(sheet.cell(row_number, column_start + offset).value)
                        for offset, field in enumerate(NAMA_FIELD_HARIAN)
                    }
                    nilai["libur_dibayar"] = (
                        nilai["jam_normal"] if hari in tanggal_libur else 0
                    )
                    nilai["jam_hadir"] = max(
                        0, nilai["jam_normal"] - nilai["libur_dibayar"]
                    )
                    rincian.append(
                        {
                            "tanggal": date(tahun, bulan, hari),
                            "id": id_karyawan,
                            "nama": nama,
                            "departemen": departemen,
                            "jabatan": jabatan,
                            "unit": nama_sheet,
                            **nilai,
                            "status": _status_harian(nilai),
                        }
                    )

        if not karyawan:
            raise ValueError("Tidak ada data karyawan yang dapat dibaca dari file audit.")

        return {
            "jenis": "audit",
            "periode": {"tahun": tahun, "bulan": bulan},
            "karyawan": karyawan,
            "rincian": rincian,
            "sheet": sheet_ditemukan,
        }
    finally:
        workbook.close()
        workbook_formula.close()
